#!/usr/bin/env python3
"""Parse PPTX files or image-slide folders/zips into a normalized JSON format.

This script has two modes:
1) --mode pptx   : full PPTX parsing using python-pptx + XML fallback logic
2) --mode images : image-per-slide import from folder or zip archive

All machine-readable output is emitted as JSON to stdout only.
All warnings/errors are emitted to stderr only.
"""

import argparse
import base64
import colorsys
import json
import os
import re
import sys
import tempfile
import traceback
import zipfile
from typing import Any, Dict, Iterable, List, Optional, Tuple

from lxml import etree

EMU_PER_INCH = 914400
DPI = 96
PT_PER_INCH = 72

NAMESPACES = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

R_EMBED_ATTR = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
R_LINK_ATTR = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link"

DEFAULT_FONT_NAME = "Calibri"
DEFAULT_FONT_SIZE_PX = 18.0
DEFAULT_FONT_COLOR = "#000000"

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


def emu_to_px(emu: Any) -> float:
    """Convert English Metric Units (EMU) to pixels using 96 DPI."""
    try:
        value = float(emu)
    except Exception:
        return 0.0
    return round((value / EMU_PER_INCH) * DPI, 4)


def pt_to_px(pt: Any) -> float:
    """Convert typographic points to pixels using 96 DPI."""
    try:
        value = float(pt)
    except Exception:
        return 0.0
    return round((value / PT_PER_INCH) * DPI, 4)


def halfpt_to_px(halfpt: Any) -> float:
    """Convert half-points to pixels.

    python-pptx can surface spacing values in half-point-like units depending on source XML.
    """
    try:
        value = float(halfpt)
    except Exception:
        return 0.0
    return pt_to_px(value / 2)


def log_warning(message: str) -> None:
    """Write a warning message to stderr only."""
    print(f"WARNING: {message}", file=sys.stderr)


def clamp(value: float, minimum: float, maximum: float) -> float:
    """Clamp a numeric value into an inclusive [minimum, maximum] range."""
    return max(minimum, min(maximum, value))


def safe_int(value: Any, default: Optional[int] = None) -> Optional[int]:
    """Safely parse an integer value and return default on failure."""
    try:
        return int(value)
    except Exception:
        return default


def safe_float(value: Any, default: Optional[float] = None) -> Optional[float]:
    """Safely parse a float value and return default on failure."""
    try:
        return float(value)
    except Exception:
        return default


def local_name(tag: str) -> str:
    """Return the XML local-name from a potentially namespaced tag."""
    if "}" in tag:
        return tag.split("}", 1)[1]
    if ":" in tag:
        return tag.split(":", 1)[1]
    return tag


def normalize_hex_color(value: Any) -> Optional[str]:
    """Normalize a color string to #RRGGBB format.

    Returns None if the value cannot be interpreted as a 6-digit RGB hex color.
    """
    if value is None:
        return None
    text = str(value).strip().replace("#", "")
    if len(text) == 3 and re.fullmatch(r"[0-9a-fA-F]{3}", text):
        text = "".join(ch * 2 for ch in text)
    if not re.fullmatch(r"[0-9a-fA-F]{6}", text):
        return None
    return f"#{text.upper()}"


def bool_from_xml_on_off(value: Any) -> Optional[bool]:
    """Convert XML on/off style values to bool when possible."""
    if value is None:
        return None
    text = str(value).strip().lower()
    if text in {"1", "true", "on", "yes"}:
        return True
    if text in {"0", "false", "off", "no"}:
        return False
    return None


def normalize_underline_value(value: Any) -> Optional[bool]:
    """Normalize underline representation to bool when possible."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"none", "0", "false", "off"}:
        return False
    return True


def map_alignment_code(value: Any) -> Optional[str]:
    """Map alignment enums/XML codes to normalized strings."""
    if value is None:
        return None

    text = getattr(value, "name", None)
    if text is None:
        text = str(value)
    normalized = text.strip().lower().replace(" ", "").replace("_", "")

    mapping = {
        "l": "left",
        "left": "left",
        "ctr": "center",
        "center": "center",
        "r": "right",
        "right": "right",
        "just": "justify",
        "justify": "justify",
        "dist": "justify",
        "thaidist": "justify",
    }
    return mapping.get(normalized)


def map_vertical_anchor(value: Any) -> str:
    """Map bodyPr anchor values to normalized vertical alignment strings."""
    if value is None:
        return "center"
    normalized = str(value).strip().lower()
    if normalized == "t":
        return "top"
    if normalized == "b":
        return "bottom"
    return "center"


def natural_sort_key(value: str) -> List[Any]:
    """Return a natural-sort key so slide2 sorts before slide10."""
    parts = re.split(r"(\d+)", value.lower())
    key: List[Any] = []
    for part in parts:
        if part.isdigit():
            key.append(int(part))
        else:
            key.append(part)
    return key


def encode_blob_to_base64(data: bytes) -> str:
    """Encode binary bytes to a base64 ASCII string."""
    return base64.b64encode(data).decode("ascii")


def apply_luminance_and_alpha(
    hex_color: str,
    lum_mod: Optional[int] = None,
    lum_off: Optional[int] = None,
    alpha: Optional[int] = None,
) -> str:
    """Apply lumMod/lumOff/alpha transforms using HLS math.

    OOXML luminance transforms are applied on lightness:
    l_new = l_old * lumMod + lumOff

    alpha is applied by compositing over black to keep #RRGGBB output.
    """
    normalized = normalize_hex_color(hex_color) or DEFAULT_FONT_COLOR
    rgb = [int(normalized[i : i + 2], 16) / 255.0 for i in (1, 3, 5)]

    h, l, s = colorsys.rgb_to_hls(rgb[0], rgb[1], rgb[2])

    mod_ratio = 1.0 if lum_mod is None else clamp(lum_mod / 100000.0, 0.0, 1.0)
    off_ratio = 0.0 if lum_off is None else clamp(lum_off / 100000.0, 0.0, 1.0)

    l = clamp((l * mod_ratio) + off_ratio, 0.0, 1.0)

    r, g, b = colorsys.hls_to_rgb(h, l, s)

    if alpha is not None:
        alpha_ratio = clamp(alpha / 100000.0, 0.0, 1.0)
        r *= alpha_ratio
        g *= alpha_ratio
        b *= alpha_ratio

    return "#%02X%02X%02X" % (
        int(round(clamp(r, 0.0, 1.0) * 255)),
        int(round(clamp(g, 0.0, 1.0) * 255)),
        int(round(clamp(b, 0.0, 1.0) * 255)),
    )


def extract_color_transforms_from_xml(color_choice: Optional[etree._Element]) -> Dict[str, Optional[int]]:
    """Extract lumMod/lumOff/alpha from an OOXML color element."""
    transforms: Dict[str, Optional[int]] = {"lum_mod": None, "lum_off": None, "alpha": None}
    if color_choice is None:
        return transforms

    for child in list(color_choice):
        name = local_name(child.tag)
        value = safe_int(child.get("val"))
        if value is None:
            continue
        if name == "lumMod":
            transforms["lum_mod"] = value
        elif name == "lumOff":
            transforms["lum_off"] = value
        elif name == "alpha":
            transforms["alpha"] = value

    return transforms


def resolve_theme_color_key_from_enum(theme_color: Any) -> Optional[str]:
    """Map python-pptx theme color enums to OOXML theme keys (dk1, accent1, etc.)."""
    if theme_color is None:
        return None

    raw = getattr(theme_color, "name", None)
    if raw is None:
        raw = str(theme_color)
    key = raw.strip().lower().replace("_", "")

    mapping = {
        "dark1": "dk1",
        "dark2": "dk2",
        "light1": "lt1",
        "light2": "lt2",
        "text1": "tx1",
        "text2": "tx2",
        "background1": "bg1",
        "background2": "bg2",
        "accent1": "accent1",
        "accent2": "accent2",
        "accent3": "accent3",
        "accent4": "accent4",
        "accent5": "accent5",
        "accent6": "accent6",
        "hyperlink": "hlink",
        "followedhyperlink": "folHlink",
        "folhlink": "folHlink",
    }

    return mapping.get(key, key)


def resolve_xml_color_choice(
    color_choice: Optional[etree._Element],
    theme_colors: Dict[str, str],
    context: str,
) -> Optional[str]:
    """Resolve an OOXML color choice element (srgbClr/schemeClr/sysClr) to #RRGGBB.

    This function is used for XML-only paths where python-pptx ColorFormat is unavailable,
    such as placeholder/list-style/default-run XML and background fills.
    """
    if color_choice is None:
        return None

    name = local_name(color_choice.tag)
    base_hex: Optional[str] = None

    if name == "srgbClr":
        base_hex = normalize_hex_color(color_choice.get("val"))
    elif name == "schemeClr":
        theme_key = (color_choice.get("val") or "").strip().lower()
        base_hex = normalize_hex_color(theme_colors.get(theme_key))
        if base_hex is None:
            log_warning(f"{context}: unresolved schemeClr '{theme_key}', using fallback #000000")
            base_hex = DEFAULT_FONT_COLOR
    elif name == "sysClr":
        base_hex = normalize_hex_color(color_choice.get("lastClr"))
        if base_hex is None:
            log_warning(f"{context}: unresolved sysClr, using fallback #000000")
            base_hex = DEFAULT_FONT_COLOR
    else:
        log_warning(f"{context}: unsupported XML color type '{name}', using fallback #000000")
        base_hex = DEFAULT_FONT_COLOR

    transforms = extract_color_transforms_from_xml(color_choice)
    return apply_luminance_and_alpha(
        base_hex,
        lum_mod=transforms["lum_mod"],
        lum_off=transforms["lum_off"],
        alpha=transforms["alpha"],
    )


def resolve_xml_solid_fill(
    solid_fill: Optional[etree._Element],
    theme_colors: Dict[str, str],
    context: str,
) -> Optional[str]:
    """Resolve an <a:solidFill> XML node to #RRGGBB."""
    if solid_fill is None:
        return None

    for child in list(solid_fill):
        # The first child color choice determines the color for solidFill.
        return resolve_xml_color_choice(child, theme_colors, context)

    return None


def extract_color_transforms_from_color_format(color_format: Any) -> Dict[str, Optional[int]]:
    """Extract lumMod/lumOff/alpha from python-pptx ColorFormat internals."""
    transforms: Dict[str, Optional[int]] = {"lum_mod": None, "lum_off": None, "alpha": None}

    try:
        # python-pptx stores underlying XML in private attributes.
        # We inspect this XML to apply OOXML transforms accurately.
        color_impl = getattr(color_format, "_color", None)
        x_color = getattr(color_impl, "_xClr", None)
        if x_color is None:
            return transforms

        for child in list(x_color):
            name = local_name(child.tag)
            value = safe_int(child.get("val"))
            if value is None:
                continue
            if name == "lumMod":
                transforms["lum_mod"] = value
            elif name == "lumOff":
                transforms["lum_off"] = value
            elif name == "alpha":
                transforms["alpha"] = value
    except Exception as exc:
        log_warning(f"extract_color_transforms_from_color_format failed: {exc}")

    return transforms


def resolve_color(color_format: Any, theme_colors: Dict[str, str]) -> Optional[str]:
    """Resolve python-pptx ColorFormat to #RRGGBB.

    Resolution order:
    - None/unknown => None
    - RGB         => direct #RRGGBB
    - THEME       => lookup from theme map + apply lumMod/lumOff/alpha
    - SYSTEM      => #000000 fallback

    Never raises; returns #000000 on failures where a color type exists but cannot resolve.
    """
    if color_format is None:
        return None

    try:
        color_type = getattr(color_format, "type", None)
        if color_type is None:
            return None

        type_name = (getattr(color_type, "name", None) or str(color_type)).upper()

        if "RGB" in type_name:
            rgb_value = getattr(color_format, "rgb", None)
            if rgb_value is None:
                log_warning("resolve_color: RGB color type with missing rgb value; using #000000")
                return DEFAULT_FONT_COLOR
            normalized = normalize_hex_color(str(rgb_value))
            if normalized is None:
                log_warning("resolve_color: invalid RGB value; using #000000")
                return DEFAULT_FONT_COLOR
            return normalized

        if "THEME" in type_name or "SCHEME" in type_name:
            theme_key = resolve_theme_color_key_from_enum(getattr(color_format, "theme_color", None))
            base_hex = normalize_hex_color(theme_colors.get((theme_key or "").lower()))
            if base_hex is None:
                log_warning(
                    f"resolve_color: unresolved theme color '{theme_key}', using #000000"
                )
                base_hex = DEFAULT_FONT_COLOR

            transforms = extract_color_transforms_from_color_format(color_format)
            return apply_luminance_and_alpha(
                base_hex,
                lum_mod=transforms["lum_mod"],
                lum_off=transforms["lum_off"],
                alpha=transforms["alpha"],
            )

        if "SYSTEM" in type_name:
            return DEFAULT_FONT_COLOR

        log_warning(f"resolve_color: unsupported color type '{type_name}', using #000000")
        return DEFAULT_FONT_COLOR
    except Exception as exc:
        log_warning(f"resolve_color failed: {exc}; using #000000")
        return DEFAULT_FONT_COLOR


def get_theme_xml_root_from_master(master: Any) -> Optional[etree._Element]:
    """Load and parse theme XML root from a slide master part.

    python-pptx exposes theme linkage via relationships, so we resolve the
    master -> theme relationship explicitly and parse the target XML blob.
    """
    theme_part = None

    try:
        master_part = master.part
    except Exception as exc:
        log_warning(f"Failed accessing slide master part: {exc}")
        return None

    try:
        theme_part = getattr(master_part, "theme_part", None)
    except Exception:
        theme_part = None

    if theme_part is None:
        try:
            from pptx.opc.constants import RELATIONSHIP_TYPE as RT

            theme_part = master_part.part_related_by(RT.THEME)
        except Exception as exc:
            log_warning(f"Failed resolving theme relationship from slide master: {exc}")
            return None

    try:
        blob = getattr(theme_part, "blob", None)
        if blob is None:
            log_warning("Theme relationship resolved but target has no blob")
            return None
        return etree.fromstring(blob)
    except Exception as exc:
        log_warning(f"Failed parsing theme XML blob: {exc}")
        return None


def build_theme_color_scheme_from_master_xml(master: Any) -> Dict[str, str]:
    """Build base theme color map (dk1, lt1, accent1..6, hlink, folHlink) from master theme XML."""
    theme_colors: Dict[str, str] = {}

    xml_root = get_theme_xml_root_from_master(master)
    if xml_root is None:
        log_warning("Failed to load theme XML from slide master")
        return theme_colors

    # We parse clrScheme directly from theme XML because this is the source of
    # canonical theme colors used by schemeClr references across slides.
    clr_scheme = xml_root.find(".//a:clrScheme", namespaces=NAMESPACES)
    if clr_scheme is None:
        log_warning("Theme XML missing a:clrScheme; theme color resolution may fallback")
        return theme_colors

    keys = [
        "dk1",
        "lt1",
        "dk2",
        "lt2",
        "accent1",
        "accent2",
        "accent3",
        "accent4",
        "accent5",
        "accent6",
        "hlink",
        "folHlink",
    ]

    for key in keys:
        entry = clr_scheme.find(f"./a:{key}", namespaces=NAMESPACES)
        if entry is None:
            log_warning(f"Theme clrScheme missing '{key}'")
            continue

        color_value: Optional[str] = None
        for child in list(entry):
            color_value = resolve_xml_color_choice(child, theme_colors, f"theme:{key}")
            if color_value:
                break

        if color_value is None:
            log_warning(f"Could not resolve theme color '{key}'")
            continue

        theme_colors[key.lower()] = color_value

    return theme_colors


def build_theme_alias_map_from_master_xml(master: Any) -> Dict[str, str]:
    """Build alias mapping (bg1/tx1/etc -> dk1/lt1/etc) from slide master clrMap."""
    alias_map: Dict[str, str] = {}

    try:
        master_xml = getattr(master.part, "_element", None)
        if master_xml is None:
            return alias_map

        clr_map = master_xml.find(".//p:clrMap", namespaces=NAMESPACES)
        if clr_map is None:
            return alias_map

        for attr_name, attr_value in clr_map.attrib.items():
            alias_key = local_name(attr_name).strip().lower()
            target_key = str(attr_value).strip().lower()
            if alias_key and target_key:
                alias_map[alias_key] = target_key
    except Exception as exc:
        log_warning(f"Failed parsing clrMap aliases from slide master XML: {exc}")

    return alias_map


def build_theme_alias_map_from_theme_color_map(master: Any) -> Dict[str, str]:
    """Build alias mapping from slide_master.theme_color_map when available.

    This complements XML clrMap parsing and keeps resolution aligned with the
    high-level python-pptx theme_color_map API.
    """
    alias_map: Dict[str, str] = {}

    try:
        theme_color_map = getattr(master, "theme_color_map", None)
        if theme_color_map is None:
            return alias_map

        attr_to_alias = {
            "background1": "bg1",
            "text1": "tx1",
            "background2": "bg2",
            "text2": "tx2",
            "accent1": "accent1",
            "accent2": "accent2",
            "accent3": "accent3",
            "accent4": "accent4",
            "accent5": "accent5",
            "accent6": "accent6",
            "hyperlink": "hlink",
            "followed_hyperlink": "folHlink",
            "followedhyperlink": "folHlink",
        }

        for attr_name, alias_key in attr_to_alias.items():
            if not hasattr(theme_color_map, attr_name):
                continue
            raw_value = getattr(theme_color_map, attr_name)
            target = resolve_theme_color_key_from_enum(raw_value)
            if target:
                alias_map[alias_key.lower()] = target.lower()
    except Exception as exc:
        log_warning(f"Failed parsing aliases from theme_color_map: {exc}")

    return alias_map


def build_theme_colors_for_slide(slide: Any) -> Dict[str, str]:
    """Build the effective theme map for a specific slide.

    We combine:
    1) canonical colors from the master theme clrScheme
    2) alias mapping from master clrMap (bg1->lt1, tx1->dk1, etc.)

    This map is then used by resolve_color for all shape/run/theme references.
    """
    theme_colors: Dict[str, str] = {}

    try:
        master = slide.slide_layout.slide_master
    except Exception as exc:
        log_warning(f"Failed to access slide master while building theme map: {exc}")
        return theme_colors

    theme_colors.update(build_theme_color_scheme_from_master_xml(master))

    alias_map = build_theme_alias_map_from_master_xml(master)
    alias_map.update(build_theme_alias_map_from_theme_color_map(master))
    for alias, target in alias_map.items():
        target_value = theme_colors.get(target.lower())
        if target_value:
            theme_colors[alias.lower()] = target_value
        else:
            log_warning(f"Theme alias '{alias}' references missing target '{target}'")

    return theme_colors


def extract_rpr_properties_from_xml(
    rpr_element: Optional[etree._Element],
    theme_colors: Dict[str, str],
    context: str,
) -> Dict[str, Any]:
    """Extract font-related properties from OOXML run properties (a:rPr/defRPr)."""
    props: Dict[str, Any] = {
        "font_name": None,
        "font_size_px": None,
        "bold": None,
        "italic": None,
        "underline": None,
        "color": None,
        "character_spacing_px": None,
    }

    if rpr_element is None:
        return props

    try:
        # a:latin/@typeface is the primary font name in DrawingML run properties.
        latin = rpr_element.find("./a:latin", namespaces=NAMESPACES)
        if latin is not None:
            props["font_name"] = latin.get("typeface")

        size_hundredth_pt = safe_int(rpr_element.get("sz"))
        if size_hundredth_pt is not None:
            props["font_size_px"] = pt_to_px(size_hundredth_pt / 100.0)

        props["bold"] = bool_from_xml_on_off(rpr_element.get("b"))
        props["italic"] = bool_from_xml_on_off(rpr_element.get("i"))
        props["underline"] = normalize_underline_value(rpr_element.get("u"))

        spc_hundredth_pt = safe_float(rpr_element.get("spc"))
        if spc_hundredth_pt is not None:
            props["character_spacing_px"] = pt_to_px(spc_hundredth_pt / 100.0)

        solid_fill = rpr_element.find("./a:solidFill", namespaces=NAMESPACES)
        if solid_fill is not None:
            props["color"] = resolve_xml_solid_fill(solid_fill, theme_colors, context)
    except Exception as exc:
        log_warning(f"{context}: failed extracting rPr properties: {exc}")

    return props


def extract_alignment_from_text_shape_xml(shape_element: Optional[etree._Element]) -> Optional[str]:
    """Extract paragraph alignment from shape XML defaults when available."""
    if shape_element is None:
        return None

    try:
        # First choice: explicit alignment on first paragraph pPr.
        ppr = shape_element.find(".//a:p/a:pPr", namespaces=NAMESPACES)
        if ppr is not None:
            alignment = map_alignment_code(ppr.get("algn"))
            if alignment is not None:
                return alignment

        # Second choice: list-style level defaults.
        lvl1 = shape_element.find(".//a:lstStyle/a:lvl1pPr", namespaces=NAMESPACES)
        if lvl1 is not None:
            return map_alignment_code(lvl1.get("algn"))
    except Exception as exc:
        log_warning(f"Failed extracting alignment from shape XML: {exc}")

    return None


def find_matching_placeholder(placeholders: Iterable[Any], idx: Any, ph_type: Any) -> Optional[Any]:
    """Find a matching placeholder by idx first, then by placeholder type."""
    by_idx: Optional[Any] = None
    by_type: Optional[Any] = None

    for placeholder in placeholders:
        try:
            ph_format = placeholder.placeholder_format
            candidate_idx = getattr(ph_format, "idx", None)
            candidate_type = getattr(ph_format, "type", None)

            if idx is not None and candidate_idx == idx:
                by_idx = placeholder
                break

            if ph_type is not None and candidate_type == ph_type and by_type is None:
                by_type = placeholder
        except Exception:
            continue

    return by_idx or by_type


def extract_placeholder_defined_props(
    placeholder: Optional[Any],
    theme_colors: Dict[str, str],
    context: str,
) -> Dict[str, Any]:
    """Extract explicitly defined text style props from a layout/master placeholder."""
    props: Dict[str, Any] = {
        "font_name": None,
        "font_size_px": None,
        "bold": None,
        "italic": None,
        "color": None,
        "alignment": None,
    }

    if placeholder is None:
        return props

    try:
        # We parse raw XML to capture defaults (lstStyle/defRPr) that python-pptx
        # high-level properties do not always surface.
        shape_element = getattr(placeholder, "element", None)
        if shape_element is None:
            return props

        props["alignment"] = extract_alignment_from_text_shape_xml(shape_element)

        # Prefer paragraph default run properties from pPr/defRPr.
        rpr = shape_element.find(".//a:p/a:pPr/a:defRPr", namespaces=NAMESPACES)
        if rpr is None:
            # Fallback to list style level defaults.
            rpr = shape_element.find(".//a:lstStyle/a:lvl1pPr/a:defRPr", namespaces=NAMESPACES)

        extracted = extract_rpr_properties_from_xml(
            rpr,
            theme_colors,
            f"{context}:placeholder-rPr",
        )

        props["font_name"] = extracted.get("font_name")
        props["font_size_px"] = extracted.get("font_size_px")
        props["bold"] = extracted.get("bold")
        props["italic"] = extracted.get("italic")
        props["color"] = extracted.get("color")
    except Exception as exc:
        log_warning(f"{context}: failed extracting placeholder-defined props: {exc}")

    return props


def resolve_placeholder_props(
    shape: Any,
    slide_layout: Any,
    slide_master: Any,
    theme_colors: Dict[str, str],
) -> Dict[str, Any]:
    """Resolve placeholder inheritance from layout/master by idx/type matching.

    This function intentionally uses XML-level style extraction from matching placeholders
    because placeholders often inherit their text defaults from layout/master definitions.
    """
    inherited: Dict[str, Any] = {
        "font_name": None,
        "font_size_px": None,
        "bold": None,
        "italic": None,
        "color": None,
        "alignment": None,
    }

    try:
        is_placeholder = getattr(shape, "is_placeholder", False)
        if not is_placeholder:
            return inherited

        ph_format = getattr(shape, "placeholder_format", None)
        if ph_format is None:
            return inherited

        ph_idx = getattr(ph_format, "idx", None)
        ph_type = getattr(ph_format, "type", None)

        layout_placeholder = find_matching_placeholder(
            getattr(slide_layout, "placeholders", []),
            ph_idx,
            ph_type,
        )
        master_placeholder = find_matching_placeholder(
            getattr(slide_master, "placeholders", []),
            ph_idx,
            ph_type,
        )

        # Inheritance precedence: master defaults first, then layout overrides.
        master_props = extract_placeholder_defined_props(
            master_placeholder,
            theme_colors,
            "master",
        )
        layout_props = extract_placeholder_defined_props(
            layout_placeholder,
            theme_colors,
            "layout",
        )

        for key in inherited.keys():
            if master_props.get(key) is not None:
                inherited[key] = master_props[key]
            if layout_props.get(key) is not None:
                inherited[key] = layout_props[key]
    except Exception as exc:
        log_warning(f"resolve_placeholder_props failed: {exc}")

    return inherited


def extract_container_text_defaults(container: Any, theme_colors: Dict[str, str], context: str) -> Dict[str, Any]:
    """Extract container-level fallback text defaults from layout or master placeholders."""
    defaults: Dict[str, Any] = {
        "font_name": None,
        "font_size_px": None,
        "bold": None,
        "italic": None,
        "color": None,
        "alignment": None,
    }

    try:
        for placeholder in getattr(container, "placeholders", []):
            props = extract_placeholder_defined_props(placeholder, theme_colors, context)
            for key, value in props.items():
                if defaults.get(key) is None and value is not None:
                    defaults[key] = value
    except Exception as exc:
        log_warning(f"{context}: failed extracting container text defaults: {exc}")

    return defaults


def choose_property(context: str, prop_name: str, candidates: List[Tuple[str, Any]], fallback: Any) -> Any:
    """Return first non-None candidate value or fallback, logging if fallback is used."""
    for source, value in candidates:
        if value is not None:
            return value

    log_warning(f"{context}: unresolved {prop_name}; using fallback {fallback!r}")
    return fallback


def extract_paragraph_bullet(ppr: Optional[etree._Element], context: str) -> Optional[Dict[str, Any]]:
    """Extract paragraph bullet metadata from pPr XML."""
    if ppr is None:
        return None

    try:
        if ppr.find("./a:buNone", namespaces=NAMESPACES) is not None:
            return None

        bu_char = ppr.find("./a:buChar", namespaces=NAMESPACES)
        if bu_char is not None:
            return {
                "type": "char",
                "char": bu_char.get("char") or "",
            }

        bu_auto = ppr.find("./a:buAutoNum", namespaces=NAMESPACES)
        if bu_auto is not None:
            return {
                "type": "auto",
                "scheme": bu_auto.get("type") or "arabicPeriod",
                "start_at": safe_int(bu_auto.get("startAt"), 1) or 1,
            }
    except Exception as exc:
        log_warning(f"{context}: failed extracting bullet properties: {exc}")

    return None


def extract_paragraph_spacing(paragraph: Any, ppr: Optional[etree._Element], which: str, context: str) -> float:
    """Extract paragraph spacing before/after in px."""
    try:
        spacing = getattr(paragraph, "space_before" if which == "before" else "space_after", None)
        if spacing is not None:
            pt_value = getattr(spacing, "pt", None)
            if pt_value is not None:
                return pt_to_px(pt_value)
            numeric = safe_float(spacing)
            if numeric is not None:
                # Requirement explicitly asks for half-point conversion helper.
                return halfpt_to_px(numeric)
    except Exception as exc:
        log_warning(f"{context}: failed reading paragraph.space_{which}: {exc}")

    if ppr is None:
        return 0.0

    try:
        tag = "spcBef" if which == "before" else "spcAft"
        spc_node = ppr.find(f"./a:{tag}", namespaces=NAMESPACES)
        if spc_node is None:
            return 0.0

        # OOXML stores explicit spacing points in a:spcPts/@val (hundredths of a point).
        spc_pts = spc_node.find("./a:spcPts", namespaces=NAMESPACES)
        if spc_pts is not None:
            value = safe_float(spc_pts.get("val"))
            if value is not None:
                return pt_to_px(value / 100.0)

        # Percentage spacing is relative and not absolute px; we keep 0 fallback.
        spc_pct = spc_node.find("./a:spcPct", namespaces=NAMESPACES)
        if spc_pct is not None:
            log_warning(f"{context}: paragraph {which} spacing is percentage-based; using 0px fallback")
    except Exception as exc:
        log_warning(f"{context}: failed parsing paragraph spacing XML: {exc}")

    return 0.0


def extract_paragraph_line_spacing(paragraph: Any, ppr: Optional[etree._Element], context: str) -> float:
    """Extract paragraph line spacing as a numeric multiplier where possible."""
    try:
        line_spacing = getattr(paragraph, "line_spacing", None)
        if isinstance(line_spacing, (int, float)):
            return round(float(line_spacing), 4)

        if line_spacing is not None:
            pt_value = getattr(line_spacing, "pt", None)
            if pt_value is not None:
                return pt_to_px(pt_value)

            numeric = safe_float(line_spacing)
            if numeric is not None:
                return round(numeric, 4)
    except Exception as exc:
        log_warning(f"{context}: failed reading paragraph line_spacing: {exc}")

    if ppr is None:
        return 1.0

    try:
        # XML fallback for line spacing values not surfaced by python-pptx.
        ln_spc = ppr.find("./a:lnSpc", namespaces=NAMESPACES)
        if ln_spc is None:
            return 1.0

        spc_pct = ln_spc.find("./a:spcPct", namespaces=NAMESPACES)
        if spc_pct is not None:
            value = safe_float(spc_pct.get("val"))
            if value is not None:
                return round(value / 100000.0, 4)

        spc_pts = ln_spc.find("./a:spcPts", namespaces=NAMESPACES)
        if spc_pts is not None:
            value = safe_float(spc_pts.get("val"))
            if value is not None:
                return pt_to_px(value / 100.0)
    except Exception as exc:
        log_warning(f"{context}: failed parsing line spacing XML: {exc}")

    return 1.0


def _length_to_px(length_value: Any) -> Optional[float]:
    """Convert python-pptx length values to px when possible."""
    if length_value is None:
        return None

    try:
        # Most python-pptx length values support .pt for typographic sizes.
        pt_value = getattr(length_value, "pt", None)
        if pt_value is not None:
            return pt_to_px(pt_value)

        emu_value = safe_float(length_value)
        if emu_value is not None:
            return emu_to_px(emu_value)
    except Exception:
        return None

    return None


def extract_run_data(
    run: Any,
    paragraph_props: Dict[str, Any],
    placeholder_props: Dict[str, Any],
    layout_defaults: Dict[str, Any],
    master_defaults: Dict[str, Any],
    theme_colors: Dict[str, str],
    context: str,
) -> Dict[str, Any]:
    """Extract a single run with full fallback resolution chain."""
    run_font = getattr(run, "font", None)
    run_text = getattr(run, "text", "")

    run_name = None
    run_size_px = None
    run_bold = None
    run_italic = None
    run_underline = None
    run_color = None
    run_char_spacing = None

    try:
        if run_font is not None:
            run_name = getattr(run_font, "name", None)
            run_size_px = _length_to_px(getattr(run_font, "size", None))
            run_bold = getattr(run_font, "bold", None)
            run_italic = getattr(run_font, "italic", None)
            run_underline = normalize_underline_value(getattr(run_font, "underline", None))
            run_color = resolve_color(getattr(run_font, "color", None), theme_colors)
    except Exception as exc:
        log_warning(f"{context}: failed reading run.font properties: {exc}")

    try:
        # XML fallback for run-only properties such as character spacing and color transforms.
        rpr = getattr(getattr(run, "_r", None), "rPr", None)
        if rpr is not None:
            rpr_props = extract_rpr_properties_from_xml(rpr, theme_colors, f"{context}:rPr")
            if run_name is None:
                run_name = rpr_props.get("font_name")
            if run_size_px is None:
                run_size_px = rpr_props.get("font_size_px")
            if run_bold is None:
                run_bold = rpr_props.get("bold")
            if run_italic is None:
                run_italic = rpr_props.get("italic")
            if run_underline is None:
                run_underline = rpr_props.get("underline")
            if run_color is None:
                run_color = rpr_props.get("color")
            run_char_spacing = rpr_props.get("character_spacing_px")
    except Exception as exc:
        log_warning(f"{context}: failed reading run XML properties: {exc}")

    font_name = choose_property(
        context,
        "font_name",
        [
            ("run", run_name),
            ("paragraph", paragraph_props.get("font_name")),
            ("placeholder", placeholder_props.get("font_name")),
            ("layout", layout_defaults.get("font_name")),
            ("master", master_defaults.get("font_name")),
        ],
        DEFAULT_FONT_NAME,
    )

    font_size_px = choose_property(
        context,
        "font_size_px",
        [
            ("run", run_size_px),
            ("paragraph", paragraph_props.get("font_size_px")),
            ("placeholder", placeholder_props.get("font_size_px")),
            ("layout", layout_defaults.get("font_size_px")),
            ("master", master_defaults.get("font_size_px")),
        ],
        DEFAULT_FONT_SIZE_PX,
    )

    bold = choose_property(
        context,
        "bold",
        [
            ("run", run_bold),
            ("paragraph", paragraph_props.get("bold")),
            ("placeholder", placeholder_props.get("bold")),
            ("layout", layout_defaults.get("bold")),
            ("master", master_defaults.get("bold")),
        ],
        False,
    )

    italic = choose_property(
        context,
        "italic",
        [
            ("run", run_italic),
            ("paragraph", paragraph_props.get("italic")),
            ("placeholder", placeholder_props.get("italic")),
            ("layout", layout_defaults.get("italic")),
            ("master", master_defaults.get("italic")),
        ],
        False,
    )

    underline = choose_property(
        context,
        "underline",
        [
            ("run", run_underline),
            ("paragraph", paragraph_props.get("underline")),
        ],
        False,
    )

    color = choose_property(
        context,
        "color",
        [
            ("run", run_color),
            ("paragraph", paragraph_props.get("color")),
            ("placeholder", placeholder_props.get("color")),
            ("layout", layout_defaults.get("color")),
            ("master", master_defaults.get("color")),
        ],
        DEFAULT_FONT_COLOR,
    )

    if run_char_spacing is None:
        run_char_spacing = choose_property(
            context,
            "character_spacing_px",
            [
                ("run", run_char_spacing),
                ("paragraph", paragraph_props.get("character_spacing_px")),
            ],
            0.0,
        )

    return {
        "text": run_text if run_text is not None else "",
        "font_name": font_name,
        "font_size_px": round(float(font_size_px), 4),
        "bold": bool(bold),
        "italic": bool(italic),
        "underline": bool(underline),
        "color": normalize_hex_color(color) or DEFAULT_FONT_COLOR,
        "character_spacing_px": round(float(run_char_spacing), 4),
    }


def extract_paragraph_data(
    paragraph: Any,
    paragraph_index: int,
    placeholder_props: Dict[str, Any],
    layout_defaults: Dict[str, Any],
    master_defaults: Dict[str, Any],
    theme_colors: Dict[str, str],
    context: str,
) -> Dict[str, Any]:
    """Extract paragraph data including spacing, bullets, and runs."""
    p_element = getattr(paragraph, "_p", None)
    ppr = None
    if p_element is not None:
        ppr = p_element.find("./a:pPr", namespaces=NAMESPACES)

    paragraph_rpr = None
    if ppr is not None:
        paragraph_rpr = ppr.find("./a:defRPr", namespaces=NAMESPACES)
    if paragraph_rpr is None and p_element is not None:
        paragraph_rpr = p_element.find("./a:endParaRPr", namespaces=NAMESPACES)

    paragraph_props = extract_rpr_properties_from_xml(
        paragraph_rpr,
        theme_colors,
        f"{context}:paragraph[{paragraph_index}]",
    )

    direct_alignment = map_alignment_code(getattr(paragraph, "alignment", None))
    xml_alignment = map_alignment_code(ppr.get("algn")) if ppr is not None else None

    alignment = choose_property(
        context,
        "paragraph_alignment",
        [
            ("paragraph", direct_alignment),
            ("paragraph_xml", xml_alignment),
            ("placeholder", placeholder_props.get("alignment")),
            ("layout", layout_defaults.get("alignment")),
            ("master", master_defaults.get("alignment")),
        ],
        "left",
    )

    space_before_px = extract_paragraph_spacing(
        paragraph,
        ppr,
        "before",
        f"{context}:paragraph[{paragraph_index}]",
    )
    space_after_px = extract_paragraph_spacing(
        paragraph,
        ppr,
        "after",
        f"{context}:paragraph[{paragraph_index}]",
    )
    line_spacing = extract_paragraph_line_spacing(
        paragraph,
        ppr,
        f"{context}:paragraph[{paragraph_index}]",
    )

    bullet = extract_paragraph_bullet(ppr, f"{context}:paragraph[{paragraph_index}]")

    runs_payload: List[Dict[str, Any]] = []
    runs = list(getattr(paragraph, "runs", []))
    if not runs:
        log_warning(f"{context}:paragraph[{paragraph_index}] has no runs; preserving empty paragraph")
        runs_payload.append(
            {
                "text": "",
                "font_name": choose_property(
                    context,
                    "font_name",
                    [
                        ("paragraph", paragraph_props.get("font_name")),
                        ("placeholder", placeholder_props.get("font_name")),
                        ("layout", layout_defaults.get("font_name")),
                        ("master", master_defaults.get("font_name")),
                    ],
                    DEFAULT_FONT_NAME,
                ),
                "font_size_px": round(
                    float(
                        choose_property(
                            context,
                            "font_size_px",
                            [
                                ("paragraph", paragraph_props.get("font_size_px")),
                                ("placeholder", placeholder_props.get("font_size_px")),
                                ("layout", layout_defaults.get("font_size_px")),
                                ("master", master_defaults.get("font_size_px")),
                            ],
                            DEFAULT_FONT_SIZE_PX,
                        )
                    ),
                    4,
                ),
                "bold": bool(
                    choose_property(
                        context,
                        "bold",
                        [
                            ("paragraph", paragraph_props.get("bold")),
                            ("placeholder", placeholder_props.get("bold")),
                            ("layout", layout_defaults.get("bold")),
                            ("master", master_defaults.get("bold")),
                        ],
                        False,
                    )
                ),
                "italic": bool(
                    choose_property(
                        context,
                        "italic",
                        [
                            ("paragraph", paragraph_props.get("italic")),
                            ("placeholder", placeholder_props.get("italic")),
                            ("layout", layout_defaults.get("italic")),
                            ("master", master_defaults.get("italic")),
                        ],
                        False,
                    )
                ),
                "underline": bool(
                    choose_property(
                        context,
                        "underline",
                        [("paragraph", paragraph_props.get("underline"))],
                        False,
                    )
                ),
                "color": normalize_hex_color(
                    choose_property(
                        context,
                        "color",
                        [
                            ("paragraph", paragraph_props.get("color")),
                            ("placeholder", placeholder_props.get("color")),
                            ("layout", layout_defaults.get("color")),
                            ("master", master_defaults.get("color")),
                        ],
                        DEFAULT_FONT_COLOR,
                    )
                )
                or DEFAULT_FONT_COLOR,
                "character_spacing_px": round(
                    float(
                        choose_property(
                            context,
                            "character_spacing_px",
                            [("paragraph", paragraph_props.get("character_spacing_px"))],
                            0.0,
                        )
                    ),
                    4,
                ),
            }
        )
    else:
        for run_index, run in enumerate(runs):
            run_context = f"{context}:paragraph[{paragraph_index}]:run[{run_index}]"
            runs_payload.append(
                extract_run_data(
                    run,
                    paragraph_props,
                    placeholder_props,
                    layout_defaults,
                    master_defaults,
                    theme_colors,
                    run_context,
                )
            )

    return {
        "alignment": alignment,
        "space_before_px": round(space_before_px, 4),
        "space_after_px": round(space_after_px, 4),
        "line_spacing": line_spacing,
        "bullet": bullet,
        "runs": runs_payload,
    }


def extract_text_body(
    text_frame: Any,
    placeholder_props: Dict[str, Any],
    layout_defaults: Dict[str, Any],
    master_defaults: Dict[str, Any],
    theme_colors: Dict[str, str],
    context: str,
) -> Dict[str, Any]:
    """Extract full text body (bodyPr + paragraphs + runs) for a text frame."""
    body_pr = getattr(getattr(text_frame, "_txBody", None), "bodyPr", None)

    # bodyPr stores text box geometry behavior (insets, wrap, vertical anchor).
    # These are needed for faithful layout in renderers.
    l_ins = safe_float(body_pr.get("lIns")) if body_pr is not None else None
    r_ins = safe_float(body_pr.get("rIns")) if body_pr is not None else None
    t_ins = safe_float(body_pr.get("tIns")) if body_pr is not None else None
    b_ins = safe_float(body_pr.get("bIns")) if body_pr is not None else None

    anchor = body_pr.get("anchor") if body_pr is not None else None
    wrap_value = body_pr.get("wrap") if body_pr is not None else None
    no_wrap_value = body_pr.get("noWrap") if body_pr is not None else None

    wrap = True
    if wrap_value is not None and str(wrap_value).strip().lower() in {"none", "false", "0"}:
        wrap = False
    if no_wrap_value is not None and str(no_wrap_value).strip().lower() in {"1", "true"}:
        wrap = False

    auto_size = getattr(text_frame, "auto_size", None)
    auto_size_name = getattr(auto_size, "name", None)
    if auto_size_name is None:
        auto_size_name = str(auto_size) if auto_size is not None else "none"

    paragraphs_payload: List[Dict[str, Any]] = []
    paragraphs = list(getattr(text_frame, "paragraphs", []))
    if not paragraphs:
        log_warning(f"{context}: text frame has no paragraphs; preserving empty paragraph list")

    for paragraph_index, paragraph in enumerate(paragraphs):
        paragraphs_payload.append(
            extract_paragraph_data(
                paragraph,
                paragraph_index,
                placeholder_props,
                layout_defaults,
                master_defaults,
                theme_colors,
                context,
            )
        )

    return {
        "inset_left_px": emu_to_px(l_ins if l_ins is not None else 0),
        "inset_right_px": emu_to_px(r_ins if r_ins is not None else 0),
        "inset_top_px": emu_to_px(t_ins if t_ins is not None else 0),
        "inset_bottom_px": emu_to_px(b_ins if b_ins is not None else 0),
        "vertical_align": map_vertical_anchor(anchor),
        "wrap": wrap,
        "autofit": auto_size_name.lower(),
        "paragraphs": paragraphs_payload,
    }


def detect_image_format_from_content_type(content_type: Optional[str], fallback_name: Optional[str] = None) -> str:
    """Detect image format string from content-type or fallback filename."""
    if content_type and "/" in content_type:
        fmt = content_type.split("/", 1)[1].strip().lower()
        if fmt in {"jpeg", "jpg"}:
            return "jpg"
        return fmt

    if fallback_name:
        ext = os.path.splitext(fallback_name)[1].lower().lstrip(".")
        if ext == "jpeg":
            return "jpg"
        if ext:
            return ext

    return "png"


def extract_image_shape_data(shape: Any, base: Dict[str, Any], context: str) -> Dict[str, Any]:
    """Extract picture shape payload with base64 image data."""
    payload = dict(base)
    payload["type"] = "image"

    try:
        image = getattr(shape, "image", None)
        if image is None:
            log_warning(f"{context}: image shape missing image blob; using empty fallback")
            payload["format"] = "png"
            payload["data"] = ""
            return payload

        blob = getattr(image, "blob", b"") or b""
        ext = getattr(image, "ext", "png") or "png"
        payload["format"] = str(ext).lower().lstrip(".")
        payload["data"] = encode_blob_to_base64(blob)
    except Exception as exc:
        log_warning(f"{context}: failed extracting image shape data: {exc}")
        payload["format"] = "png"
        payload["data"] = ""

    return payload


def extract_cell_fill_color(cell: Any, theme_colors: Dict[str, str], context: str) -> Optional[str]:
    """Extract table cell fill color from API first, XML fallback second."""
    try:
        fill = getattr(cell, "fill", None)
        if fill is not None and getattr(fill, "type", None) is not None:
            fill_type_name = (getattr(fill.type, "name", None) or str(fill.type)).upper()
            if "SOLID" in fill_type_name:
                return resolve_color(getattr(fill, "fore_color", None), theme_colors)
    except Exception as exc:
        log_warning(f"{context}: failed reading cell fill from API: {exc}")

    try:
        tc_pr = getattr(getattr(cell, "_tc", None), "tcPr", None)
        if tc_pr is not None:
            solid = tc_pr.find("./a:solidFill", namespaces=NAMESPACES)
            if solid is not None:
                return resolve_xml_solid_fill(solid, theme_colors, f"{context}:tcPr:solidFill")
    except Exception as exc:
        log_warning(f"{context}: failed reading cell fill from XML: {exc}")

    return None


def extract_cell_borders(cell: Any, theme_colors: Dict[str, str], context: str) -> Dict[str, Optional[Dict[str, Any]]]:
    """Extract table cell border colors/widths from tcPr XML."""
    borders: Dict[str, Optional[Dict[str, Any]]] = {
        "left": None,
        "right": None,
        "top": None,
        "bottom": None,
    }

    side_map = {
        "left": "lnL",
        "right": "lnR",
        "top": "lnT",
        "bottom": "lnB",
    }

    try:
        tc_pr = getattr(getattr(cell, "_tc", None), "tcPr", None)
        if tc_pr is None:
            return borders

        for side_name, xml_name in side_map.items():
            ln = tc_pr.find(f"./a:{xml_name}", namespaces=NAMESPACES)
            if ln is None:
                continue

            width = emu_to_px(safe_float(ln.get("w"), 0.0) or 0.0)
            color = None
            solid = ln.find("./a:solidFill", namespaces=NAMESPACES)
            if solid is not None:
                color = resolve_xml_solid_fill(solid, theme_colors, f"{context}:{side_name}")

            borders[side_name] = {
                "color": normalize_hex_color(color) if color else None,
                "width_px": width,
            }
    except Exception as exc:
        log_warning(f"{context}: failed extracting table cell borders: {exc}")

    return borders


def extract_table_cell_data(
    cell: Any,
    row_index: int,
    col_index: int,
    placeholder_props: Dict[str, Any],
    layout_defaults: Dict[str, Any],
    master_defaults: Dict[str, Any],
    theme_colors: Dict[str, str],
    context: str,
) -> Dict[str, Any]:
    """Extract full table cell payload including text, fill, borders, and merge spans."""
    cell_context = f"{context}:cell[{row_index},{col_index}]"

    text_body = extract_text_body(
        cell.text_frame,
        placeholder_props,
        layout_defaults,
        master_defaults,
        theme_colors,
        cell_context,
    )

    row_span = 1
    col_span = 1

    try:
        is_merge_origin = bool(getattr(cell, "is_merge_origin", False))
        is_spanned = bool(getattr(cell, "is_spanned", False))

        if is_merge_origin:
            row_span = safe_int(getattr(cell, "span_height", 1), 1) or 1
            col_span = safe_int(getattr(cell, "span_width", 1), 1) or 1
        elif is_spanned:
            # Spanned cells are covered by a merge origin. Keep explicit 0 spans for clarity.
            row_span = 0
            col_span = 0
    except Exception as exc:
        log_warning(f"{cell_context}: failed resolving merge spans: {exc}")

    return {
        "row_index": row_index,
        "col_index": col_index,
        "row_span": row_span,
        "col_span": col_span,
        "fill_color": normalize_hex_color(extract_cell_fill_color(cell, theme_colors, cell_context)),
        "borders": extract_cell_borders(cell, theme_colors, cell_context),
        "text_body": text_body,
    }


def extract_table_shape_data(
    shape: Any,
    base: Dict[str, Any],
    placeholder_props: Dict[str, Any],
    layout_defaults: Dict[str, Any],
    master_defaults: Dict[str, Any],
    theme_colors: Dict[str, str],
    context: str,
) -> Dict[str, Any]:
    """Extract table shape including row/column metrics and per-cell content/styles."""
    payload = dict(base)
    payload["type"] = "table"

    rows_payload: List[Dict[str, Any]] = []
    columns_payload: List[Dict[str, Any]] = []

    try:
        table = getattr(shape, "table", None)
        if table is None:
            log_warning(f"{context}: table shape has no table object; using empty table fallback")
            payload["columns"] = []
            payload["rows"] = []
            return payload

        for col_index, column in enumerate(getattr(table, "columns", [])):
            col_width_px = emu_to_px(getattr(column, "width", 0))
            columns_payload.append({"index": col_index, "width_px": col_width_px})

        for row_index, row in enumerate(getattr(table, "rows", [])):
            row_height_px = emu_to_px(getattr(row, "height", 0))
            row_cells: List[Dict[str, Any]] = []
            for col_index, cell in enumerate(getattr(row, "cells", [])):
                row_cells.append(
                    extract_table_cell_data(
                        cell,
                        row_index,
                        col_index,
                        placeholder_props,
                        layout_defaults,
                        master_defaults,
                        theme_colors,
                        context,
                    )
                )

            rows_payload.append(
                {
                    "index": row_index,
                    "height_px": row_height_px,
                    "cells": row_cells,
                }
            )
    except Exception as exc:
        log_warning(f"{context}: failed extracting table shape data: {exc}")

    payload["columns"] = columns_payload
    payload["rows"] = rows_payload
    return payload


def extract_plain_shape_data(shape: Any, base: Dict[str, Any], theme_colors: Dict[str, str], context: str) -> Dict[str, Any]:
    """Extract non-text geometric shape fill/line styling."""
    payload = dict(base)
    payload["type"] = "shape"

    fill_color = None
    line_color = None
    line_width_px = 0.0

    try:
        fill = getattr(shape, "fill", None)
        if fill is not None and getattr(fill, "type", None) is not None:
            fill_type_name = (getattr(fill.type, "name", None) or str(fill.type)).upper()
            if "SOLID" in fill_type_name:
                fill_color = resolve_color(getattr(fill, "fore_color", None), theme_colors)
    except Exception as exc:
        log_warning(f"{context}: failed extracting shape fill color: {exc}")

    try:
        line = getattr(shape, "line", None)
        if line is not None:
            line_color = resolve_color(getattr(line, "color", None), theme_colors)
            line_width_px = _length_to_px(getattr(line, "width", None)) or 0.0
    except Exception as exc:
        log_warning(f"{context}: failed extracting shape line style: {exc}")

    auto_shape_type = getattr(shape, "auto_shape_type", None)
    auto_shape_name = getattr(auto_shape_type, "name", None) or str(auto_shape_type)

    payload["shape_type_name"] = auto_shape_name
    payload["fill_color"] = normalize_hex_color(fill_color) if fill_color else None
    payload["line_color"] = normalize_hex_color(line_color) if line_color else None
    payload["line_width_px"] = round(float(line_width_px), 4)

    return payload


def extract_line_shape_data(shape: Any, base: Dict[str, Any], theme_colors: Dict[str, str], context: str) -> Dict[str, Any]:
    """Extract line shape styling details."""
    payload = dict(base)
    payload["type"] = "line"

    line_color = None
    line_width_px = 0.0

    try:
        line = getattr(shape, "line", None)
        if line is not None:
            line_color = resolve_color(getattr(line, "color", None), theme_colors)
            line_width_px = _length_to_px(getattr(line, "width", None)) or 0.0
    except Exception as exc:
        log_warning(f"{context}: failed extracting line shape style: {exc}")

    payload["line_color"] = normalize_hex_color(line_color) if line_color else None
    payload["line_width_px"] = round(float(line_width_px), 4)
    return payload


def classify_shape_type(shape: Any) -> str:
    """Classify a shape into one of: text/image/shape/table/group/line/unknown."""
    try:
        shape_type_name = (getattr(getattr(shape, "shape_type", None), "name", None) or str(getattr(shape, "shape_type", ""))).upper()
    except Exception:
        shape_type_name = ""

    if "GROUP" in shape_type_name:
        return "group"

    if bool(getattr(shape, "has_table", False)):
        return "table"

    if "PICTURE" in shape_type_name:
        return "image"

    if bool(getattr(shape, "has_text_frame", False)):
        return "text"

    if "LINE" in shape_type_name:
        return "line"

    if "AUTO" in shape_type_name or "FREEFORM" in shape_type_name or "CALLOUT" in shape_type_name:
        return "shape"

    return "unknown"


def get_part_xml_root(part: Any) -> Optional[etree._Element]:
    """Return XML root element for a pptx part when available."""
    root = getattr(part, "_element", None)
    if root is not None:
        return root
    root = getattr(part, "element", None)
    if root is not None:
        return root
    return None


def get_related_part_bytes(part: Any, rel_id: str, context: str) -> Tuple[Optional[bytes], Optional[str], Optional[str]]:
    """Resolve a relationship ID to bytes/content-type/name from a pptx part."""
    try:
        related = part.related_part(rel_id)
        blob = getattr(related, "blob", None)
        content_type = getattr(related, "content_type", None)
        partname = str(getattr(related, "partname", ""))
        if blob is None:
            log_warning(f"{context}: related part '{rel_id}' has no blob")
            return None, content_type, partname
        return blob, content_type, partname
    except Exception as exc:
        log_warning(f"{context}: failed resolving related part '{rel_id}': {exc}")
        return None, None, None


def extract_background_from_part(part: Any, theme_colors: Dict[str, str], context: str) -> Optional[Dict[str, Any]]:
    """Extract background definition from a slide/layout/master part.

    XML-level parsing is required because inherited backgrounds and blip fills are
    not consistently surfaced by python-pptx high-level API.
    """
    root = get_part_xml_root(part)
    if root is None:
        return None

    try:
        # A part background is stored in p:cSld/p:bg when explicitly defined.
        bg = root.find(".//p:cSld/p:bg", namespaces=NAMESPACES)
        if bg is None:
            return None

        bg_pr = bg.find("./p:bgPr", namespaces=NAMESPACES)
        bg_ref = bg.find("./p:bgRef", namespaces=NAMESPACES)

        if bg_pr is not None:
            solid = bg_pr.find("./a:solidFill", namespaces=NAMESPACES)
            if solid is not None:
                color = resolve_xml_solid_fill(solid, theme_colors, f"{context}:bgPr:solidFill")
                if color is None:
                    log_warning(f"{context}: background solidFill unresolved, using #000000")
                    color = DEFAULT_FONT_COLOR
                return {"type": "solid", "color": normalize_hex_color(color) or DEFAULT_FONT_COLOR}

            blip_fill = bg_pr.find("./a:blipFill", namespaces=NAMESPACES)
            if blip_fill is not None:
                blip = blip_fill.find("./a:blip", namespaces=NAMESPACES)
                if blip is not None:
                    rel_id = blip.get(R_EMBED_ATTR) or blip.get(R_LINK_ATTR)
                    if rel_id:
                        blob, content_type, partname = get_related_part_bytes(part, rel_id, f"{context}:bgPr:blipFill")
                        if blob is not None:
                            fmt = detect_image_format_from_content_type(content_type, partname)
                            return {
                                "type": "image",
                                "data": encode_blob_to_base64(blob),
                                "format": fmt,
                            }
                    else:
                        log_warning(f"{context}: background blipFill missing relationship id")

        if bg_ref is not None:
            # bgRef can reference scheme colors from style matrices.
            # We try to resolve direct color child for a safe solid fallback.
            for child in list(bg_ref):
                color = resolve_xml_color_choice(child, theme_colors, f"{context}:bgRef")
                if color:
                    return {"type": "solid", "color": normalize_hex_color(color) or DEFAULT_FONT_COLOR}

    except Exception as exc:
        log_warning(f"{context}: failed extracting background: {exc}")

    return None


def resolve_slide_background(slide: Any, theme_colors: Dict[str, str]) -> Dict[str, Any]:
    """Resolve effective slide background with inheritance fallback.

    Precedence:
    1) explicit slide background
    2) layout background
    3) master background
    4) solid black fallback
    """
    try:
        slide_bg = extract_background_from_part(slide.part, theme_colors, "slide")
        if slide_bg is not None:
            return slide_bg
    except Exception as exc:
        log_warning(f"slide background extraction failed: {exc}")

    try:
        layout = slide.slide_layout
        layout_bg = extract_background_from_part(layout.part, theme_colors, "layout")
        if layout_bg is not None:
            return layout_bg
    except Exception as exc:
        log_warning(f"layout background extraction failed: {exc}")

    try:
        master = slide.slide_layout.slide_master
        master_bg = extract_background_from_part(master.part, theme_colors, "master")
        if master_bg is not None:
            return master_bg
    except Exception as exc:
        log_warning(f"master background extraction failed: {exc}")

    log_warning("background unresolved after slide/layout/master checks; using #000000")
    return {"type": "solid", "color": DEFAULT_FONT_COLOR}


def extract_shape_data(
    shape: Any,
    slide: Any,
    theme_colors: Dict[str, str],
    layout_defaults: Dict[str, Any],
    master_defaults: Dict[str, Any],
    offset_x_px: float,
    offset_y_px: float,
    context: str,
) -> Dict[str, Any]:
    """Extract one shape object and normalize all geometric values to pixels."""
    shape_id = getattr(shape, "shape_id", None)
    shape_name = getattr(shape, "name", "")

    x_px = round(emu_to_px(getattr(shape, "left", 0)) + offset_x_px, 4)
    y_px = round(emu_to_px(getattr(shape, "top", 0)) + offset_y_px, 4)
    width_px = round(emu_to_px(getattr(shape, "width", 0)), 4)
    height_px = round(emu_to_px(getattr(shape, "height", 0)), 4)

    rotation = safe_float(getattr(shape, "rotation", 0.0), 0.0) or 0.0

    base = {
        "id": shape_id,
        "name": shape_name,
        "x": x_px,
        "y": y_px,
        "width": width_px,
        "height": height_px,
        "rotation": round(rotation, 4),
    }

    shape_kind = classify_shape_type(shape)
    placeholder_props = resolve_placeholder_props(
        shape,
        slide.slide_layout,
        slide.slide_layout.slide_master,
        theme_colors,
    )

    if shape_kind == "text":
        payload = dict(base)
        payload["type"] = "text"

        try:
            text_frame = getattr(shape, "text_frame", None)
            if text_frame is None:
                log_warning(f"{context}: text shape missing text_frame; using empty fallback")
                payload["text_body"] = {
                    "inset_left_px": 0.0,
                    "inset_right_px": 0.0,
                    "inset_top_px": 0.0,
                    "inset_bottom_px": 0.0,
                    "vertical_align": "center",
                    "wrap": True,
                    "autofit": "none",
                    "paragraphs": [],
                }
                return payload

            payload["text_body"] = extract_text_body(
                text_frame,
                placeholder_props,
                layout_defaults,
                master_defaults,
                theme_colors,
                context,
            )
        except Exception as exc:
            log_warning(f"{context}: failed extracting text shape; using safe fallback: {exc}")
            payload["text_body"] = {
                "inset_left_px": 0.0,
                "inset_right_px": 0.0,
                "inset_top_px": 0.0,
                "inset_bottom_px": 0.0,
                "vertical_align": "center",
                "wrap": True,
                "autofit": "none",
                "paragraphs": [],
            }

        return payload

    if shape_kind == "image":
        return extract_image_shape_data(shape, base, context)

    if shape_kind == "table":
        return extract_table_shape_data(
            shape,
            base,
            placeholder_props,
            layout_defaults,
            master_defaults,
            theme_colors,
            context,
        )

    if shape_kind == "line":
        return extract_line_shape_data(shape, base, theme_colors, context)

    if shape_kind == "shape":
        return extract_plain_shape_data(shape, base, theme_colors, context)

    if shape_kind == "group":
        payload = dict(base)
        payload["type"] = "group"
        children: List[Dict[str, Any]] = []

        try:
            # Child shape coordinates are resolved relative to the group and then
            # offset by the group's absolute x/y, per requirement.
            for child_index, child in enumerate(getattr(shape, "shapes", [])):
                child_context = f"{context}:group_child[{child_index}]"
                children.append(
                    extract_shape_data(
                        child,
                        slide,
                        theme_colors,
                        layout_defaults,
                        master_defaults,
                        x_px,
                        y_px,
                        child_context,
                    )
                )
        except Exception as exc:
            log_warning(f"{context}: failed extracting group children: {exc}")

        payload["children"] = children
        return payload

    shape_type_name = getattr(getattr(shape, "shape_type", None), "name", None) or str(getattr(shape, "shape_type", ""))
    log_warning(f"{context}: unrecognized shape type '{shape_type_name}', exporting as unknown")

    payload = dict(base)
    payload["type"] = "unknown"
    payload["shape_type_name"] = shape_type_name
    return payload


def extract_shapes_for_slide(
    slide: Any,
    theme_colors: Dict[str, str],
    slide_index: int,
) -> List[Dict[str, Any]]:
    """Extract all shapes for a single slide."""
    shapes_payload: List[Dict[str, Any]] = []

    layout_defaults = extract_container_text_defaults(slide.slide_layout, theme_colors, f"slide[{slide_index}]:layout")
    master_defaults = extract_container_text_defaults(
        slide.slide_layout.slide_master,
        theme_colors,
        f"slide[{slide_index}]:master",
    )

    for shape_index, shape in enumerate(getattr(slide, "shapes", [])):
        context = f"slide[{slide_index}]:shape[{shape_index}]"
        try:
            shapes_payload.append(
                extract_shape_data(
                    shape,
                    slide,
                    theme_colors,
                    layout_defaults,
                    master_defaults,
                    0.0,
                    0.0,
                    context,
                )
            )
        except Exception as exc:
            log_warning(f"{context}: shape extraction failed completely, using unknown fallback: {exc}")
            shapes_payload.append(
                {
                    "id": getattr(shape, "shape_id", None),
                    "name": getattr(shape, "name", ""),
                    "type": "unknown",
                    "x": round(emu_to_px(getattr(shape, "left", 0)), 4),
                    "y": round(emu_to_px(getattr(shape, "top", 0)), 4),
                    "width": round(emu_to_px(getattr(shape, "width", 0)), 4),
                    "height": round(emu_to_px(getattr(shape, "height", 0)), 4),
                    "rotation": round(safe_float(getattr(shape, "rotation", 0.0), 0.0) or 0.0, 4),
                }
            )

    return shapes_payload


def parse_pptx_mode(pptx_path: str) -> Dict[str, Any]:
    """Parse PPTX file into normalized JSON payload."""
    try:
        from pptx import Presentation
    except Exception as exc:
        raise RuntimeError(
            "python-pptx is required for --mode pptx. Install with: pip install python-pptx"
        ) from exc

    prs = Presentation(pptx_path)

    presentation_payload: Dict[str, Any] = {
        "mode": "pptx",
        "slide_width_px": emu_to_px(prs.slide_width),
        "slide_height_px": emu_to_px(prs.slide_height),
        "slides": [],
    }

    for index, slide in enumerate(prs.slides):
        theme_colors = build_theme_colors_for_slide(slide)

        slide_payload = {
            "index": index,
            "background": resolve_slide_background(slide, theme_colors),
            "shapes": extract_shapes_for_slide(slide, theme_colors, index),
        }
        presentation_payload["slides"].append(slide_payload)

    return presentation_payload


def collect_image_files_from_folder(folder_path: str) -> List[str]:
    """Collect supported image files recursively from a folder."""
    files: List[str] = []
    for root, _dirs, names in os.walk(folder_path):
        for name in names:
            ext = os.path.splitext(name)[1].lower()
            if ext in IMAGE_EXTENSIONS:
                files.append(os.path.join(root, name))

    files.sort(key=lambda item: (natural_sort_key(os.path.basename(item)), item.lower()))
    return files

def parse_images_mode(path_value: str) -> Dict[str, Any]:
    """Parse image-slide input from a folder or zip archive."""
    try:
        from PIL import Image
    except Exception as exc:
        raise RuntimeError(
            "Pillow is required for --mode images. Install with: pip install Pillow"
        ) from exc

    slides: List[Dict[str, Any]] = []

    if os.path.isdir(path_value):
        image_files = collect_image_files_from_folder(path_value)

        for index, image_path in enumerate(image_files):
            try:
                with open(image_path, "rb") as file_handle:
                    raw_bytes = file_handle.read()

                with Image.open(image_path) as image_obj:
                    width_px, height_px = image_obj.size

                ext = os.path.splitext(image_path)[1].lower().lstrip(".") or "png"
                if ext == "jpeg":
                    ext = "jpg"

                slides.append(
                    {
                        "index": index,
                        "width_px": int(width_px),
                        "height_px": int(height_px),
                        "format": ext,
                        "data": encode_blob_to_base64(raw_bytes),
                    }
                )
            except Exception as exc:
                log_warning(f"images mode: failed reading image '{image_path}': {exc}")

    elif zipfile.is_zipfile(path_value):
        with tempfile.TemporaryDirectory(prefix="parse_pptx_images_") as temp_dir:
            with zipfile.ZipFile(path_value, "r") as archive:
                archive.extractall(temp_dir)

            image_files = collect_image_files_from_folder(temp_dir)

            for index, image_path in enumerate(image_files):
                try:
                    with open(image_path, "rb") as file_handle:
                        raw_bytes = file_handle.read()

                    with Image.open(image_path) as image_obj:
                        width_px, height_px = image_obj.size

                    ext = os.path.splitext(image_path)[1].lower().lstrip(".") or "png"
                    if ext == "jpeg":
                        ext = "jpg"

                    slides.append(
                        {
                            "index": index,
                            "width_px": int(width_px),
                            "height_px": int(height_px),
                            "format": ext,
                            "data": encode_blob_to_base64(raw_bytes),
                        }
                    )
                except Exception as exc:
                    log_warning(f"images mode: failed reading image '{image_path}': {exc}")
    else:
        raise ValueError("--mode images expects a folder path or a zip file path")

    if not slides:
        log_warning("images mode: no supported image files found; returning empty slides list")
        slide_width_px = 0
        slide_height_px = 0
    else:
        slide_width_px = slides[0]["width_px"]
        slide_height_px = slides[0]["height_px"]

    return {
        "mode": "images",
        "slide_width_px": slide_width_px,
        "slide_height_px": slide_height_px,
        "slides": slides,
    }


def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(description="Parse PPTX or image slides into JSON")
    parser.add_argument("--mode", required=True, choices=["pptx", "images"], help="Import mode")
    parser.add_argument("path", help="Path to .pptx file or images folder/zip")
    return parser.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    """Program entrypoint. Prints JSON to stdout and returns process exit code."""
    args = parse_args(argv)

    if args.mode == "pptx":
        payload = parse_pptx_mode(args.path)
    else:
        payload = parse_images_mode(args.path)

    # stdout MUST contain only valid JSON.
    sys.stdout.write(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception:
        traceback.print_exc(file=sys.stderr)
        raise SystemExit(1)
