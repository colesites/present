#!/usr/bin/env python3
"""Build the PPTX fixtures the importer is tested against.

    python3 scripts/make-pptx-fixtures.py        # needs python-pptx

Each fixture exercises what real services' decks use — placeholders that
inherit from their layout, theme colours through the master's colour map,
groups, tables, lines, pictures, backgrounds — plus the edge cases the parser
has rules for. `node scripts/pptx-expected.mjs` then records what the Electron
importer (parse_pptx.py and lib/pptx-payload.mjs) makes of each, and the Rust
importer is required to make exactly the same.

Everything is generated: no real presentation is committed.
"""

import copy
import io
import os
import struct
import zlib

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = os.path.join(os.path.dirname(__file__), "..", "tests", "fixtures", "pptx")
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def qn(tag):
    prefix, name = tag.split(":")
    return "{%s}%s" % ({"a": A, "p": P, "r": R}[prefix], name)


def png(width, height, rgb):
    """A tiny solid PNG, built by hand so no image library is needed."""
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(kind, data):
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def encoded(fmt, size, rgb):
    """A solid image in `fmt`, through Pillow, which python-pptx needs anyway."""
    from PIL import Image

    stream = io.BytesIO()
    Image.new("RGB", size, rgb).save(stream, format=fmt)
    return stream.getvalue()


def layout(prs, name):
    return next(item for item in prs.slide_layouts if item.name == name)


def add_run(paragraph, text, **font):
    run = paragraph.add_run()
    run.text = text
    for key, value in font.items():
        if key == "color":
            run.font.color.rgb = RGBColor.from_string(value)
        else:
            setattr(run.font, key, value)
    return run


def set_scheme_color(run, key, **transforms):
    """Give a run a theme colour, with lumMod/lumOff/alpha children."""
    rpr = run._r.get_or_add_rPr()
    for fill in rpr.findall(qn("a:solidFill")):
        rpr.remove(fill)
    solid = etree.SubElement(rpr, qn("a:solidFill"))
    scheme = etree.SubElement(solid, qn("a:schemeClr"), val=key)
    for name, value in transforms.items():
        etree.SubElement(scheme, qn("a:" + name), val=str(value))
    # solidFill must precede latin in CT_TextCharacterProperties.
    latin = rpr.find(qn("a:latin"))
    if latin is not None:
        rpr.remove(latin)
        rpr.append(latin)


def text_and_theme(path):
    prs = Presentation()

    # Placeholders inheriting position and styles from the Title Slide layout.
    slide = prs.slides.add_slide(layout(prs, "Title Slide"))
    slide.shapes.title.text = "Amazing Grace"
    subtitle = slide.placeholders[1]
    subtitle.text_frame.text = "How sweet the sound"
    subtitle.text_frame.add_paragraph().text = "That saved a wretch like me"

    # Runs, fonts and paragraph properties in a text box.
    slide = prs.slides.add_slide(layout(prs, "Blank"))
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(5))
    frame = box.text_frame
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    frame.margin_left = Emu(182880)
    frame.margin_top = Emu(45720)

    first = frame.paragraphs[0]
    first.alignment = PP_ALIGN.CENTER
    add_run(first, "Bold ", bold=True, size=Pt(40), name="Georgia")
    add_run(first, "italic ", italic=True, color="C00000")
    add_run(first, "underlined", underline=True)
    first.space_before = Pt(12)
    first.space_after = Pt(6)
    first.line_spacing = 1.5

    second = frame.add_paragraph()
    second.alignment = PP_ALIGN.RIGHT
    second.line_spacing = Pt(30)
    themed = add_run(second, "accent, darkened ")
    set_scheme_color(themed, "accent1", lumMod=75000)
    lightened = add_run(second, "text, lightened")
    set_scheme_color(lightened, "tx1", lumMod=50000, lumOff=50000)
    spaced = add_run(second, " spaced")
    spaced._r.get_or_add_rPr().set("spc", "300")
    doubled = add_run(second, " double underline")
    doubled._r.get_or_add_rPr().set("u", "dbl")

    third = frame.add_paragraph()
    third._p.get_or_add_pPr().set("algn", "dist")
    add_run(third, "distributed")
    spc = etree.SubElement(third._p.get_or_add_pPr(), qn("a:spcBef"))
    etree.SubElement(spc, qn("a:spcPct"), val="50000")

    # An empty paragraph, sized only by its end-of-paragraph properties.
    empty = frame.add_paragraph()
    end = etree.SubElement(empty._p, qn("a:endParaRPr"), lang="en-US", sz="2800", b="1")
    fill = etree.SubElement(end, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr"), val="00B050")

    # A line break inside a paragraph: the runs either side, and no break text.
    broken = frame.add_paragraph()
    add_run(broken, "before")
    broken._p.append(etree.Element(qn("a:br")))
    add_run(broken, "after")

    # Body placeholder on Title and Content, text only — styles from the master.
    slide = prs.slides.add_slide(layout(prs, "Title and Content"))
    slide.shapes.title.text = "Verse 1"
    body = slide.placeholders[1].text_frame
    body.text = "Through many dangers, toils and snares"
    body.add_paragraph().text = "I have already come"
    body.paragraphs[1].level = 1

    prs.save(path)


def shapes_tables_groups(path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(layout(prs, "Blank"))

    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1.5))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    rect.line.color.rgb = RGBColor(0xFF, 0xC0, 0x00)
    rect.line.width = Pt(2.25)
    rect.rotation = 15.5
    rect.text_frame.text = "In a shape"

    # An auto shape with no text body at all.
    bare = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5), Inches(1), Inches(1), Inches(1))
    txbody = bare._element.find(qn("p:txBody"))
    if txbody is not None:
        bare._element.remove(txbody)

    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(3), Inches(6), Inches(3.5))
    connector.line.color.rgb = RGBColor(0x70, 0x30, 0xA0)
    connector.line.width = Pt(3)
    plain_connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(4), Inches(2), Inches(4))

    group = slide.shapes.add_group_shape()
    child_a = group.shapes.add_textbox(Inches(7), Inches(1), Inches(2), Inches(1))
    child_a.text_frame.text = "grouped A"
    inner = group.shapes.add_group_shape()
    child_b = inner.shapes.add_textbox(Inches(8), Inches(2.5), Inches(2), Inches(1))
    child_b.text_frame.text = "nested B"

    table_shape = slide.shapes.add_table(3, 3, Inches(1), Inches(5), Inches(6), Inches(1.5))
    table = table_shape.table
    for row in range(3):
        for col in range(3):
            table.cell(row, col).text = f"r{row}c{col}"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(2, 2).text = ""

    prs.save(path)


def pictures_and_backgrounds(path):
    prs = Presentation()

    slide = prs.slides.add_slide(layout(prs, "Blank"))
    slide.shapes.add_picture(io.BytesIO(png(4, 3, (255, 0, 0))), Inches(1), Inches(1), Inches(2), Inches(1.5))
    slide.shapes.add_picture(io.BytesIO(encoded("JPEG", (3, 2), (250, 200, 0))), Inches(4), Inches(1), Inches(1), Inches(1))
    slide.shapes.add_picture(io.BytesIO(encoded("GIF", (2, 2), (255, 0, 0))), Inches(6), Inches(1), Inches(1), Inches(1))
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(0x10, 0x20, 0x30)

    # A background picture, which python-pptx has no API for.
    slide = prs.slides.add_slide(layout(prs, "Blank"))
    _, rel_id = slide.part.get_or_add_image_part(io.BytesIO(png(8, 6, (0, 128, 255))))
    c_sld = slide._element.find(qn("p:cSld"))
    bg = etree.Element(qn("p:bg"))
    bg_pr = etree.SubElement(bg, qn("p:bgPr"))
    blip_fill = etree.SubElement(bg_pr, qn("a:blipFill"))
    etree.SubElement(blip_fill, qn("a:blip")).set(qn("r:embed"), rel_id)
    etree.SubElement(etree.SubElement(blip_fill, qn("a:stretch")), qn("a:fillRect"))
    etree.SubElement(bg_pr, qn("a:effectLst"))
    c_sld.insert(0, bg)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "over a picture"

    # A theme-coloured background with a transform.
    slide = prs.slides.add_slide(layout(prs, "Blank"))
    c_sld = slide._element.find(qn("p:cSld"))
    bg = etree.Element(qn("p:bg"))
    bg_pr = etree.SubElement(bg, qn("p:bgPr"))
    solid = etree.SubElement(bg_pr, qn("a:solidFill"))
    scheme = etree.SubElement(solid, qn("a:schemeClr"), val="accent2")
    etree.SubElement(scheme, qn("a:lumMod"), val="60000")
    etree.SubElement(scheme, qn("a:lumOff"), val="40000")
    etree.SubElement(bg_pr, qn("a:effectLst"))
    c_sld.insert(0, bg)

    # A picture placeholder holding a picture: not a picture to the importer.
    slide = prs.slides.add_slide(layout(prs, "Picture with Caption"))
    slide.placeholders[1].insert_picture(io.BytesIO(png(4, 4, (0, 200, 0))))
    slide.shapes.title.text = "Captioned"

    prs.save(path)


def empty(path):
    Presentation().save(path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for name, build in [
        ("text-and-theme", text_and_theme),
        ("shapes-tables-groups", shapes_tables_groups),
        ("pictures-and-backgrounds", pictures_and_backgrounds),
        ("empty", empty),
    ]:
        target = os.path.join(OUT, name + ".pptx")
        build(target)
        print("wrote", os.path.relpath(target))
